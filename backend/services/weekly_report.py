"""班级周报服务（家校闭环）：按班级聚合窗口期学情，产出报告与导出。

口径（窗口默认近 7 天，含今天）：
- 新增错题 / 复习次数 / 复习正确率（good+easy 占比）
- 当前待复习：due_at 已到且未归档（归档口径与 QuestionOut.mastered 一致：reps≥3 且间隔≥21 天）
- 薄弱知识点：标签掌握度 < 50% 的前 3 个（复用看板 build_tag_stats 口径）

AI 只读数不造数：报告所有数字都来自 questions / review_logs 的真实聚合。
"""
from __future__ import annotations

import datetime as dt
import io
from collections.abc import Callable, Iterator
from contextlib import contextmanager

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from sqlalchemy import select
from sqlalchemy.orm import Session, sessionmaker

from backend.models.orm import ClassMember, Question, ReviewLog, SchoolClass, User
from backend.services.stats import build_tag_stats
from backend.utils.logging import get_logger

logger = get_logger(__name__)

WEAK_TAG_THRESHOLD = 0.5
_WEAK_TAG_LIMIT = 3
_GOOD_GRADES = {"good", "easy"}


class ClassAccessDenied(LookupError):
    """班级不存在或不属于该教师（对外一律按 404，不泄露存在性）。"""


def _mastered(question: Question) -> bool:
    return question.reps >= 3 and question.interval_days >= 21


def _as_utc(value: dt.datetime | None) -> dt.datetime | None:
    if value is None:
        return None
    return value if value.tzinfo else value.replace(tzinfo=dt.timezone.utc)


class WeeklyReportService:
    def __init__(
        self,
        session_factory: sessionmaker | Callable[[], Iterator[Session]] | None = None,
    ):
        self._session_factory = session_factory

    @contextmanager
    def _session(self) -> Iterator[Session]:
        from backend.database import session_scope

        with session_scope(self._session_factory) as session:
            yield session

    def build(
        self,
        teacher_id: int,
        class_id: int,
        *,
        days: int = 7,
        today: dt.date | None = None,
    ) -> dict:
        """生成班级学情报告（JSON 可序列化，供 API 与前端共用）。"""
        days = max(1, min(int(days), 31))
        today = today or dt.date.today()
        period_start = today - dt.timedelta(days=days - 1)
        window_start = dt.datetime(
            period_start.year, period_start.month, period_start.day, tzinfo=dt.timezone.utc
        )
        window_end = dt.datetime(
            today.year, today.month, today.day, tzinfo=dt.timezone.utc
        ) + dt.timedelta(days=1)
        now = dt.datetime.now(dt.timezone.utc)

        with self._session() as session:
            klass = session.get(SchoolClass, class_id)
            if klass is None or klass.teacher_id != teacher_id:
                raise ClassAccessDenied(
                    f"班级不可访问 class_id={class_id} teacher_id={teacher_id}"
                )
            members = session.execute(
                select(User.id, User.username)
                .join(ClassMember, ClassMember.student_id == User.id)
                .where(ClassMember.class_id == class_id)
                .order_by(User.username.asc())
            ).all()

            students: list[dict] = []
            for student_id, username in members:
                questions = (
                    session.execute(
                        select(Question).where(Question.user_id == student_id)
                    )
                    .scalars()
                    .all()
                )
                logs = (
                    session.execute(
                        select(ReviewLog).where(ReviewLog.user_id == student_id)
                    )
                    .scalars()
                    .all()
                )
                students.append(
                    self._student_row(
                        student_id, username, questions, logs,
                        window_start, window_end, now,
                    )
                )

        students.sort(
            key=lambda r: (-(r["reviews"] + r["created"]), r["username"])
        )
        return {
            "class_id": class_id,
            "class_name": klass.name,
            "period": {
                "start": period_start.isoformat(),
                "end": today.isoformat(),
                "days": days,
            },
            "generated_at": now.isoformat(),
            "students": students,
            "summary": {
                "students": len(students),
                "created": sum(r["created"] for r in students),
                "reviews": sum(r["reviews"] for r in students),
                "active": sum(
                    1 for r in students if r["created"] > 0 or r["reviews"] > 0
                ),
            },
        }

    def _student_row(
        self,
        student_id: int,
        username: str,
        questions: list[Question],
        logs: list[ReviewLog],
        window_start: dt.datetime,
        window_end: dt.datetime,
        now: dt.datetime,
    ) -> dict:
        created = [
            q
            for q in questions
            if window_start <= (_as_utc(q.created_at) or window_end) < window_end
        ]
        reviewed = [
            log
            for log in logs
            if window_start <= (_as_utc(log.reviewed_at) or window_end) < window_end
        ]
        strong = sum(1 for log in reviewed if log.grade in _GOOD_GRADES)
        overdue = sum(
            1
            for q in questions
            if not _mastered(q)
            and _as_utc(q.due_at) is not None
            and _as_utc(q.due_at) <= now
        )
        return {
            "user_id": student_id,
            "username": username,
            "total": len(questions),
            "created": len(created),
            "reviews": len(reviewed),
            "accuracy": round(strong / len(reviewed) * 100) if reviewed else None,
            "overdue": overdue,
            "weak_tags": self._weak_tags(questions, logs),
        }

    @staticmethod
    def _weak_tags(questions: list[Question], logs: list[ReviewLog]) -> list[dict]:
        logs_by_question: dict[int, list[tuple[str, float]]] = {}
        for log in logs:
            logs_by_question.setdefault(log.question_id, []).append(
                (log.grade, log.next_interval)
            )
        stats = build_tag_stats(questions, logs_by_question)
        weak = [s for s in stats if s.count > 0 and s.mastery < WEAK_TAG_THRESHOLD]
        weak.sort(key=lambda s: (s.mastery, -s.count))
        return [
            {"tag": s.tag, "mastery": s.mastery, "count": s.count}
            for s in weak[:_WEAK_TAG_LIMIT]
        ]


def render_markdown(report: dict) -> str:
    """报告 → Markdown（前端预览 / 家长群粘贴）。"""
    period = report["period"]
    summary = report["summary"]
    lines = [
        f"## 📣 {report['class_name']} · 学情周报",
        f"**统计窗口**：{period['start']} ~ {period['end']}（{period['days']} 天）　"
        f"**学生**：{summary['students']} 人（活跃 {summary['active']}）　"
        f"**新增错题**：{summary['created']}　**复习**：{summary['reviews']} 次",
        "",
        "| 学生 | 新增错题 | 复习次数 | 正确率 | 待复习 | 薄弱知识点 |",
        "| --- | ---: | ---: | ---: | ---: | --- |",
    ]
    for row in report["students"]:
        accuracy = f"{row['accuracy']}%" if row["accuracy"] is not None else "—"
        weak = (
            "、".join(
                f"{w['tag']}({int(w['mastery'] * 100)}%)" for w in row["weak_tags"]
            )
            or "—"
        )
        lines.append(
            f"| {row['username']} | {row['created']} | {row['reviews']} "
            f"| {accuracy} | {row['overdue']} | {weak} |"
        )
    return "\n".join(lines)


def generate_word_report(report: dict) -> io.BytesIO:
    """报告 → Word 周报（教师可直接下发家长）。"""
    doc = Document()
    heading = doc.add_heading(f"{report['class_name']} · 学情周报", level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    period = report["period"]
    summary = report["summary"]
    subtitle = doc.add_paragraph(
        f"统计窗口 {period['start']} ~ {period['end']}（{period['days']} 天） · "
        f"学生 {summary['students']} 人（活跃 {summary['active']}） · "
        f"MathMaster Edu 生成 · {dt.date.today():%Y-%m-%d}"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(10)

    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Grid Accent 1"
    header = table.rows[0].cells
    for i, text in enumerate(
        ["学生", "新增错题", "复习次数", "正确率", "待复习", "薄弱知识点"]
    ):
        header[i].text = text
        for paragraph in header[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True

    for row in report["students"]:
        cells = table.add_row().cells
        accuracy = f"{row['accuracy']}%" if row["accuracy"] is not None else "—"
        weak = (
            "、".join(
                f"{w['tag']}({int(w['mastery'] * 100)}%)" for w in row["weak_tags"]
            )
            or "—"
        )
        cells[0].text = row["username"]
        cells[1].text = str(row["created"])
        cells[2].text = str(row["reviews"])
        cells[3].text = accuracy
        cells[4].text = str(row["overdue"])
        cells[5].text = weak

    doc.add_paragraph()
    note = doc.add_paragraph(
        "说明：待复习为当前已到期未归档的错题数；薄弱知识点为掌握度低于 50% 的标签。"
        "数据来自学生真实录入与复习记录。"
    )
    note.runs[0].font.size = Pt(9)

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream
