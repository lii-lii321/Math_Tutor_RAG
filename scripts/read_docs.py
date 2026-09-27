"""读取用户提供的两份 docx 文档（一次性脚本）。"""

from docx import Document

for name, path in [
    ("提升分析", r"D:\个人下载\Desktop\提升分析.docx"),
    ("具体手册", r"D:\个人下载\Desktop\具体手册.docx"),
]:
    print(f"===== {name} =====")
    try:
        doc = Document(path)
        for para in doc.paragraphs:
            if para.text.strip():
                print(para.text)
        for table in doc.tables:
            print("--- 表格 ---")
            for row in table.rows:
                print(" | ".join(cell.text.strip() for cell in row.cells))
    except Exception as exc:  # noqa: BLE001
        print(f"读取失败: {exc}")
    print()
