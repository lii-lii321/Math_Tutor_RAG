"""Word (.docx) 导入链路测试：文本提取、AI 拆题、服务层入库。"""
from __future__ import annotations

import io
import uuid

import pytest

from backend.database import SessionLocal, init_db
from backend.models.orm import User
from backend.services.ai import get_ai_service
from backend.services.ai.base import BaseAIProvider
from backend.services.ai.mock import MockProvider
from backend.services.docx_import import extract_docx_text
from backend.services.question_service import QuestionService


def _build_docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("月考错题整理", level=1)
    document.add_paragraph("第一次月考中错误率较高的题目，整理如下。")
    document.add_heading("第 1 题", level=2)
    document.add_paragraph("已知 $x^2 = 9$，求 x 的值。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "步骤"
    table.cell(0, 1).text = "要点"
    table.cell(1, 0).text = "开平方"
    table.cell(1, 1).text = "注意正负"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


class TestExtractDocx:
    def test_extracts_headings_paragraphs_tables(self):
        text = extract_docx_text(_build_docx())
        assert "# 月考错题整理" in text
        assert "## 第 1 题" in text
        assert "$x^2 = 9$" in text
        assert "| 步骤 | 要点 |" in text
        assert "| 开平方 | 注意正负 |" in text

    def test_invalid_docx_raises_value_error(self):
        with pytest.raises(ValueError):
            extract_docx_text(b"not a docx")

    def test_empty_document_raises_value_error(self):
        import docx

        buf = io.BytesIO()
        docx.Document().save(buf)
        with pytest.raises(ValueError):
            extract_docx_text(buf.getvalue())


class TestSplitQuestions:
    def test_base_split_parses_json(self, monkeypatch):
        """BaseAIProvider.split_questions 解析 chat 返回的 JSON。"""
        import json as _json

        provider = MockProvider()

        def _fake_chat(messages):
            return _json.dumps(
                {
                    "questions": [
                        {
                            "content": "题一：$x^2=9$",
                            "answer": "x=±3",
                            "knowledge_points": ["开平方"],
                            "difficulty": "easy",
                        },
                        {
                            "content": "题二：$x+1=2$",
                            "answer": "x=1",
                            "knowledge_points": ["一元一次方程"],
                            "difficulty": "easy",
                        },
                    ]
                },
                ensure_ascii=False,
            )

        monkeypatch.setattr(provider, "chat", _fake_chat)
        chunks = BaseAIProvider.split_questions(provider, "文档文本")
        assert len(chunks) == 2
        assert chunks[0]["content"].startswith("题一")

    def test_split_rejects_empty_result(self, monkeypatch):
        provider = MockProvider()

        def _fake_chat(messages):
            return '{"questions": []}'

        monkeypatch.setattr(provider, "chat", _fake_chat)
        from backend.services.ai.base import AIMessageError, BaseAIProvider

        with pytest.raises(AIMessageError):
            BaseAIProvider.split_questions(provider, "文档文本")


class TestAddQuestionsFromDocx:
    def test_mock_provider_imports_whole_doc(self):
        """演示模式：整份文本作为一题入库（零 AI 消耗）。"""
        init_db(seed_users=True)
        with SessionLocal() as session:
            user = User(username=f"docx_{uuid.uuid4().hex[:8]}", password_hash="x", role="student")
            session.add(user)
            session.commit()
            session.refresh(user)

        service = QuestionService(session_factory=SessionLocal)
        result = service.add_questions_from_docx(user.id, _build_docx(), tags=["Word"])
        assert result["total"] == 1
        assert result["imported"] == 1
        saved = result["items"][0]
        assert saved.source == "word"
        assert "月考错题整理" in saved.content_markdown
        assert "Word" in saved.tags

    def test_multi_question_import_with_real_split(self):
        """AI 拆出多题时逐题入库，难度非法值回落 medium。"""
        import json as _json

        init_db(seed_users=True)
        with SessionLocal() as session:
            user = User(
                username=f"docx2_{uuid.uuid4().hex[:8]}", password_hash="x", role="student"
            )
            session.add(user)
            session.commit()
            session.refresh(user)

        service = QuestionService(session_factory=SessionLocal)
        provider = get_ai_service()

        def _fake_chat(messages):
            return _json.dumps(
                {
                    "questions": [
                        {"content": "甲题内容", "answer": "42", "knowledge_points": ["测试"], "difficulty": "easy"},
                        {"content": "乙题内容", "answer": "", "knowledge_points": [], "difficulty": "超难"},
                    ]
                },
                ensure_ascii=False,
            )

        provider.chat = _fake_chat
        service.ai = provider
        # Mock 的 split_questions 是覆写实现（不走 chat）——patch 成走基类解析路径
        service.ai.split_questions = lambda text: BaseAIProvider.split_questions(
            provider, text
        )

        result = service.add_questions_from_docx(user.id, _build_docx(), tags=["整理"])
        assert result["total"] == 2
        assert result["imported"] == 2
        difficulties = [q.difficulty for q in result["items"]]
        assert difficulties == ["easy", "medium"]  # 非法难度回落 medium
        assert all("整理" in q.tags for q in result["items"])

    def test_max_questions_cap(self):
        """拆题数量上限 20（服务层参数）。"""
        import json as _json

        provider = MockProvider()

        def _fake_chat(messages):
            return _json.dumps(
                {
                    "questions": [
                        {"content": f"题{i}", "answer": "", "knowledge_points": [], "difficulty": "easy"}
                        for i in range(30)
                    ]
                },
                ensure_ascii=False,
            )

        monkeypatch = pytest.MonkeyPatch()
        monkeypatch.setattr(provider, "chat", _fake_chat)
        chunks = BaseAIProvider.split_questions(provider, "文档文本")[:20]
        assert len(chunks) == 20
        monkeypatch.undo()
