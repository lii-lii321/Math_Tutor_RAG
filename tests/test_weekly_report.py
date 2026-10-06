"""班级周报测试：聚合口径、越权防护、Markdown/Word 导出、API；个人周报（E3）。"""
from __future__ import annotations

import datetime as dt
import uuid

import pytest

from backend.database import SessionLocal, init_db
from backend.models.orm import Question, ReviewLog, User
from backend.services.class_service import ClassService
from backend.services.question_service import QuestionService
from backend.services.weekly_report import (
    ClassAccessDenied,
    WeeklyReportService,
    build_for_user,
    generate_word_report,
    generate_word_report_self,
    render_markdown,
    render_markdown_self,
)


@pytest.fixture
def teacher() -> User:
    init_db(seed_users=True)
    with SessionLocal() as session:
        user = User(username=f"t_{uuid.uuid4().hex[:8]}", password_hash="x", role="teacher")
        session.add(user)
        session.commit()
        session.refresh(user)
        return user


@pytest.fixture
def student() -> User:
    with SessionLocal() as session:
        user = User(username=f"s_{uuid.uuid4().hex[:8]}", password_hash="x", role="student")
        session.add(user)
        session.commit()
        session.refresh(user)
        return user


def _seed_review(user_id: int, question_id: int, grade: str = "good") -> None:
    with SessionLocal() as session:
        session.add(
            ReviewLog(
                question_id=question_id,
                user_id=user_id,
                grade=grade,
                quality=5 if grade in ("good", "easy") else 0,
                prev_interval=0,
                next_interval=1,
                ease_after=2.5,
            )
        )
        session.commit()


class TestWeeklyReportService:
    def test_build_report_full_row(self, teacher, student):
        class_service = ClassService()
        klass = class_service.create_class(teacher.id, "周报班")
        class_service.add_student(teacher.id, klass["id"], student.id)

        question_service = QuestionService(session_factory=SessionLocal)
        question_service.create_manual_question(
            student.id, content_markdown="分数题", tags=["分数"], answer="x"
        )
        question_service.create_manual_question(
            student.id, content_markdown="方程题", tags=["方程"], answer="y"
        )
        with SessionLocal() as session:
            question = session.query(Question).filter_by(user_id=student.id).first()
            question_id, user_id = question.id, student.id
        _seed_review(user_id, question_id, grade="again")  # 复习失败 → 弱点
        _seed_review(user_id, question_id, grade="good")

        # 一道已到期未归档的题（due_at 置为过去）
        with SessionLocal() as session:
            overdue_question = session.query(Question).filter_by(user_id=student.id).all()[1]
            overdue_question.due_at = dt.datetime.now(dt.timezone.utc) - dt.timedelta(days=1)
            session.commit()

        report = WeeklyReportService().build(teacher.id, klass["id"], days=7)

        assert report["class_name"] == "周报班"
        assert report["period"]["days"] == 7
        assert report["summary"]["students"] == 1
        row = report["students"][0]
        assert row["username"] == student.username
        assert row["created"] == 2
        assert row["reviews"] == 2
        assert row["accuracy"] == 50  # again + good
        assert row["overdue"] == 1
        assert any(w["tag"] in ("分数", "方程") for w in row["weak_tags"])

    def test_inactive_student_zeros(self, teacher, student):
        class_service = ClassService()
        klass = class_service.create_class(teacher.id, "安静班")
        class_service.add_student(teacher.id, klass["id"], student.id)

        report = WeeklyReportService().build(teacher.id, klass["id"], days=7)
        row = report["students"][0]
        assert row["created"] == 0 and row["reviews"] == 0
        assert row["accuracy"] is None
        assert row["weak_tags"] == []
        assert report["summary"]["active"] == 0

    def test_other_teacher_denied(self, teacher, student):
        class_service = ClassService()
        klass = class_service.create_class(teacher.id, "别人的班")
        with SessionLocal() as session:
            stranger = User(
                username=f"t2_{uuid.uuid4().hex[:8]}", password_hash="x", role="teacher"
            )
            session.add(stranger)
            session.commit()
            stranger_id = stranger.id
        with pytest.raises(ClassAccessDenied):
            WeeklyReportService().build(stranger_id, klass["id"])

    def test_days_clamped(self, teacher, student):
        class_service = ClassService()
        klass = class_service.create_class(teacher.id, "窗口班")
        report = WeeklyReportService().build(teacher.id, klass["id"], days=999)
        assert report["period"]["days"] == 31


class TestReportRendering:
    def _report(self, teacher_id: int, class_id: int, username: str) -> dict:
        return {
            "class_id": class_id,
            "class_name": "渲染班",
            "period": {"start": "2026-09-25", "end": "2026-10-01", "days": 7},
            "generated_at": "2026-10-01T00:00:00+00:00",
            "students": [
                {
                    "user_id": 1,
                    "username": username,
                    "total": 3,
                    "created": 2,
                    "reviews": 5,
                    "accuracy": 80,
                    "overdue": 1,
                    "weak_tags": [{"tag": "分数", "mastery": 0.2, "count": 2}],
                }
            ],
            "summary": {"students": 1, "created": 2, "reviews": 5, "active": 1},
        }

    def test_markdown_contains_row(self, teacher, student):
        report = self._report(teacher.id, 1, student.username)
        text = render_markdown(report)
        assert student.username in text
        assert "渲染班" in text
        assert "分数(20%)" in text

    def test_word_export_opens(self, teacher, student):
        from io import BytesIO

        from docx import Document

        stream = generate_word_report(self._report(teacher.id, 1, student.username))
        doc = Document(BytesIO(stream.getvalue()))
        table = doc.tables[0]
        assert table.rows[0].cells[0].text == "学生"
        assert any(student.username in row.cells[0].text for row in table.rows)


class TestWeeklyReportAPI:
    @pytest.fixture
    def teacher_headers(self, client):
        username = f"wr_t_{uuid.uuid4().hex[:8]}"
        client.post(
            "/api/auth/register",
            json={
                "username": username,
                "password": "secret1",
                "role": "teacher",
                "invite_code": "test-invite-code",
            },
        )
        login = client.post("/api/auth/login", json={"username": username, "password": "secret1"})
        return {"Authorization": f"Bearer {login.json()['access_token']}"}

    @pytest.fixture
    def student_auth(self, client):
        username = f"wr_s_{uuid.uuid4().hex[:8]}"
        client.post(
            "/api/auth/register",
            json={"username": username, "password": "secret1", "role": "student"},
        )
        login = client.post("/api/auth/login", json={"username": username, "password": "secret1"})
        body = login.json()
        return {
            "headers": {"Authorization": f"Bearer {body['access_token']}"},
            "user_id": body["user_id"],
        }

    def test_report_endpoint(self, client, teacher_headers, student_auth):
        created = client.post("/api/classes", headers=teacher_headers, json={"name": "API 周报班"})
        class_id = created.json()["id"]
        client.post(
            f"/api/classes/{class_id}/members",
            headers=teacher_headers,
            json={"student_id": student_auth["user_id"]},
        )

        response = client.get(
            f"/api/classes/{class_id}/weekly-report?days=7", headers=teacher_headers
        )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["class_name"] == "API 周报班"
        assert body["students"][0]["username"].startswith("wr_s_")

    def test_days_validation(self, client, teacher_headers, student_auth):
        created = client.post("/api/classes", headers=teacher_headers, json={"name": "窗口班"})
        class_id = created.json()["id"]
        bad = client.get(
            f"/api/classes/{class_id}/weekly-report?days=0", headers=teacher_headers
        )
        assert bad.status_code == 422

    def test_other_teacher_class_404(self, client, teacher_headers, student_auth):
        created = client.post("/api/classes", headers=teacher_headers, json={"name": "私有班"})
        class_id = created.json()["id"]
        username = f"wr_t2_{uuid.uuid4().hex[:8]}"
        client.post(
            "/api/auth/register",
            json={
                "username": username,
                "password": "secret1",
                "role": "teacher",
                "invite_code": "test-invite-code",
            },
        )
        login = client.post("/api/auth/login", json={"username": username, "password": "secret1"})
        stranger = {"Authorization": f"Bearer {login.json()['access_token']}"}

        response = client.get(f"/api/classes/{class_id}/weekly-report", headers=stranger)
        assert response.status_code == 404

    def test_student_forbidden(self, client, student_auth):
        response = client.get(
            "/api/classes/1/weekly-report", headers=student_auth["headers"]
        )
        assert response.status_code == 403

    def test_requires_auth(self, client):
        assert client.get("/api/classes/1/weekly-report").status_code == 401


class TestSelfReport:
    """个人周报（E3）：窗口过滤、口径、隔离、夹取与导出。"""

    def _backdate(self, *, question=None, log=None, days: int) -> None:
        shift = dt.timedelta(days=days)
        with SessionLocal() as session:
            if question is not None:
                row = session.get(Question, question.id)
                row.created_at = dt.datetime.now(dt.timezone.utc) - shift
            if log is not None:
                row = session.get(ReviewLog, log.id)
                row.reviewed_at = dt.datetime.now(dt.timezone.utc) - shift
            session.commit()

    def test_window_filter_and_accuracy(self, student):
        """窗口外不计；正确率 = good+easy 占比。"""
        service = QuestionService(session_factory=SessionLocal)
        old_q = service.create_manual_question(
            student.id, content_markdown="窗口外题", answer="a"
        )
        self._backdate(question=old_q, days=40)
        _seed_review(student.id, old_q.id, grade="good")
        with SessionLocal() as session:
            old_log = (
                session.query(ReviewLog)
                .filter_by(user_id=student.id)
                .order_by(ReviewLog.id.desc())
                .first()
            )
        self._backdate(log=old_log, days=40)
        new_q = service.create_manual_question(
            student.id, content_markdown="窗口内题", answer="b"
        )
        for grade in ("good", "easy", "again"):
            _seed_review(student.id, new_q.id, grade=grade)

        report = build_for_user(student.id, days=7)
        assert report["title"] == f"{student.username} · 学习周报"
        assert report["period"]["days"] == 7
        row = report["row"]
        assert row["created"] == 1  # 窗口外题不计
        assert row["reviews"] == 3  # 窗口外日志不计
        assert row["accuracy"] == 67  # (good+easy)/3

    def test_overdue_excludes_archived(self, student):
        """overdue 排除已归档题（reps≥3 且 interval≥21）。"""
        service = QuestionService(session_factory=SessionLocal)
        archived = service.create_manual_question(
            student.id, content_markdown="已归档题", answer="a"
        )
        fresh = service.create_manual_question(
            student.id, content_markdown="未归档到期题", answer="b"
        )
        past = dt.datetime.now(dt.timezone.utc) - dt.timedelta(days=1)
        with SessionLocal() as session:
            a = session.get(Question, archived.id)
            a.reps, a.interval_days, a.due_at = 3, 21, past
            f = session.get(Question, fresh.id)
            f.due_at = past
            session.commit()

        row = build_for_user(student.id)["row"]
        assert row["overdue"] == 1  # 仅未归档的到期题

    def test_user_isolation(self, student):
        other = User(username=f"s2_{uuid.uuid4().hex[:8]}", password_hash="x", role="student")
        with SessionLocal() as session:
            session.add(other)
            session.commit()
            other_id = other.id
        service = QuestionService(session_factory=SessionLocal)
        service.create_manual_question(
            student.id, content_markdown="只属于 A 的题", answer="x"
        )

        row_a = build_for_user(student.id)["row"]
        row_b = build_for_user(other_id)["row"]
        assert row_a["total"] == 1 and row_a["created"] == 1
        assert row_b["total"] == 0 and row_b["created"] == 0 and row_b["reviews"] == 0

    def test_days_clamped(self, student):
        assert build_for_user(student.id, days=999)["period"]["days"] == 31
        assert build_for_user(student.id, days=0)["period"]["days"] == 1

    def test_render_markdown_self_contains_username_and_numbers(self, student):
        report = {
            "title": f"{student.username} · 学习周报",
            "period": {"start": "2026-09-25", "end": "2026-10-01", "days": 7},
            "generated_at": "2026-10-01T00:00:00+00:00",
            "row": {
                "user_id": student.id,
                "username": student.username,
                "total": 12,
                "created": 3,
                "reviews": 9,
                "accuracy": 78,
                "overdue": 2,
                "weak_tags": [{"tag": "函数", "mastery": 0.3, "count": 2}],
            },
        }
        text = render_markdown_self(report)
        assert student.username in text
        assert "📘" in text and "学习周报" in text
        assert "| 3 | 9 | 78% | 2 | 函数(30%) |" in text
        assert "累计错题**：12" in text

    def test_word_report_self_reopenable(self, student):
        from io import BytesIO

        from docx import Document

        report = {
            "title": f"{student.username} · 学习周报",
            "period": {"start": "2026-09-25", "end": "2026-10-01", "days": 7},
            "generated_at": "2026-10-01T00:00:00+00:00",
            "row": {
                "user_id": student.id,
                "username": student.username,
                "total": 12,
                "created": 3,
                "reviews": 9,
                "accuracy": 78,
                "overdue": 2,
                "weak_tags": [],
            },
        }
        stream = generate_word_report_self(report)
        doc = Document(BytesIO(stream.getvalue()))
        assert student.username in doc.paragraphs[0].text
        table = doc.tables[0]
        assert len(table.columns) == 5  # 无「学生」列
        assert table.rows[0].cells[0].text == "新增错题"
        assert table.rows[1].cells[2].text == "78%"

    def test_empty_data_accuracy_none_no_zero_division(self, student):
        """空数据：accuracy None 不零除，渲染为「—」。"""
        report = build_for_user(student.id)
        assert report["row"]["accuracy"] is None
        assert report["row"]["created"] == 0
        text = render_markdown_self(report)
        assert "—" in text
        stream = generate_word_report_self(report)
        assert stream.getvalue()[:2] == b"PK"  # 有效 docx（zip）
